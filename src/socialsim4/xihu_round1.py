from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import shutil
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from socialsim4.backend.services.documents import extract_text_from_pdf


DEFAULT_XIHU_PACKAGE_ID = "round1"
ROOT_DIR = Path(__file__).resolve().parents[2]
DATA_DIR = Path(os.getenv("SOCIALSIM4_DATA_DIR") or ROOT_DIR / "data")
XIHU_PACKAGE_ROOT = DATA_DIR / "xihu_packages"
ARM_IDS = [f"A{index}" for index in range(9)]
ARM_LABELS = {
    "A0": "A0 控制组",
    "A1": "A1 官方复杂材料",
    "A2": "A2 简明图文",
    "A3": "A3 家庭+个人收益",
    "A4": "A4 个人损失",
    "A5": "A5 家庭损失",
    "A6": "A6 个人收益",
    "A7": "A7 竞品负面+官方正面",
    "A8": "A8 视频信息",
}
ARM_FRAME_SUMMARIES = {
    "A0": "控制组不额外提供干预材料，参与者主要依赖既有认知与周围讨论作出判断。",
    "A1": "官方复杂材料强调完整政策条款、价格、保障责任与官方表达，信息密度较高。",
    "A2": "简明图文材料用更易读的结构浓缩保障对象、保费、报销范围与服务亮点。",
    "A3": "家庭+个人收益材料通过正向案例同时强调个人报销收益与家庭风险缓释。",
    "A4": "个人损失材料用损失框架强调个人未投保时可能承担的医疗开支。",
    "A5": "家庭损失材料用损失框架强调未投保如何把医疗负担转移给整个家庭。",
    "A6": "个人收益材料用收益框架突出个人投保后可能获得的经济回报与保障价值。",
    "A7": "负面信息材料先引入竞品负面信息，再叠加政府监管与官方正面信息来重塑判断。",
    "A8": "视频信息材料用更短、更叙事化的视频形式压缩核心政策信息，降低理解门槛。",
}
MATERIAL_KIND_LABELS = {
    "intervention_pdf": "干预材料",
    "questionnaire_doc": "问卷设计",
    "video": "视频材料",
    "analysis_doc": "研究复核材料",
    "analysis_script": "分析脚本",
}
NUMERIC_PATTERN = re.compile(r"^-?\d+(?:\.\d+)?$")
WHITESPACE_PATTERN = re.compile(r"\s+")
NON_WORD_PATTERN = re.compile(r"[^0-9a-zA-Z\u4e00-\u9fff]+")
XIHU_METRIC_DEFINITIONS = [
    {
        "key": "self_enrollment_willingness",
        "label": "为自己投保意愿",
        "patterns": [["为自己"], ["自己", "投保"], ["自己", "参保"], ["自己", "可能性"]],
    },
    {
        "key": "parents_enrollment_willingness",
        "label": "为父母投保意愿",
        "patterns": [["为父母"], ["父母", "投保"], ["父母", "参保"], ["父母", "可能性"]],
    },
    {
        "key": "children_enrollment_willingness",
        "label": "为子女投保意愿",
        "patterns": [["为子女"], ["子女", "投保"], ["子女", "参保"], ["子女", "可能性"]],
    },
    {
        "key": "spouse_enrollment_willingness",
        "label": "为配偶投保意愿",
        "patterns": [["为配偶"], ["配偶", "投保"], ["配偶", "参保"], ["配偶", "可能性"]],
    },
    {
        "key": "recommendation_willingness",
        "label": "推荐意愿",
        "patterns": [["推荐", "亲朋好友"], ["推荐", "西湖益联保"], ["推荐", "可能性"]],
    },
    {
        "key": "next_year_enrollment_willingness",
        "label": "下一年参保意愿",
        "patterns": [["下一年度", "参保"], ["下一年", "参保"], ["下一年度", "投保"], ["下一年", "投保"]],
    },
    {
        "key": "information_comprehension",
        "label": "信息理解度",
        "patterns": [["读懂"], ["理解", "信息"], ["理解", "以上信息"]],
    },
    {
        "key": "personal_impact_perception",
        "label": "个人影响感知",
        "patterns": [["个人生活", "影响"], ["对您个人", "影响"], ["个人", "影响程度"]],
    },
    {
        "key": "family_impact_perception",
        "label": "家庭影响感知",
        "patterns": [["家庭生活", "影响"], ["对您家庭", "影响"], ["家庭", "影响程度"]],
    },
    {
        "key": "gain_frame_perception",
        "label": "收益价值感知",
        "patterns": [["价值", "多大"], ["财务保护"], ["收益", "认可"]],
    },
    {
        "key": "loss_frame_perception",
        "label": "损失风险感知",
        "patterns": [["不需要额外", "保险保障"], ["损失", "感知"], ["医疗负担", "担心"]],
    },
    {
        "key": "government_endorsement",
        "label": "政府背书感知",
        "patterns": [["政府", "监管"], ["政府", "有力监管"], ["发布来源", "权威性"]],
    },
    {
        "key": "medical_burden_protection",
        "label": "医疗负担保护感知",
        "patterns": [["必要的财务保护"], ["医疗负担"], ["权益保护"]],
    },
]


def ensure_xihu_package_root() -> Path:
    XIHU_PACKAGE_ROOT.mkdir(parents=True, exist_ok=True)
    return XIHU_PACKAGE_ROOT


def get_xihu_package_dir(package_id: str) -> Path:
    return ensure_xihu_package_root() / package_id


def get_xihu_package_path(package_id: str) -> Path:
    return get_xihu_package_dir(package_id) / "package.json"


def _normalize_text(value: Any) -> str:
    text = str(value or "")
    text = text.replace("（", "(").replace("）", ")").replace("＆", "&")
    text = WHITESPACE_PATTERN.sub("", text)
    return text.lower()


def _slugify(value: str) -> str:
    normalized = NON_WORD_PATTERN.sub("-", value.lower()).strip("-")
    return normalized or "item"


def _truncate(value: str, limit: int = 220) -> str:
    compact = WHITESPACE_PATTERN.sub(" ", value).strip()
    if len(compact) <= limit:
        return compact
    return f"{compact[:limit]}..."


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    digest.update(path.read_bytes())
    return digest.hexdigest()


def _visible_files(directory: Path) -> list[Path]:
    return sorted(
        [
            path
            for path in directory.iterdir()
            if path.is_file() and not path.name.startswith("~$") and not path.name.startswith(".~")
        ],
        key=lambda item: item.name,
    )


def _match_single_path(root: Path, prefix: str) -> Path:
    matches = sorted(root.glob(f"{prefix}-*"))
    if len(matches) != 1:
        raise ValueError(f"Expected exactly one directory for {prefix}, found {len(matches)}")
    return matches[0]


def _material_kind_for_file(arm_id: str | None, path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return "intervention_pdf"
    if suffix == ".docx":
        if "问卷设计" in path.name:
            return "questionnaire_doc"
        if arm_id:
            return "intervention_pdf"
        return "analysis_doc"
    if suffix == ".mp4":
        return "video"
    if suffix == ".do":
        return "analysis_script"
    return "analysis_doc"


def _extract_material_text(path: Path) -> str:
    suffix = path.suffix.lower()
    content = path.read_bytes()
    if suffix == ".pdf":
        return extract_text_from_pdf(content)
    if suffix == ".docx":
        import docx

        document = docx.Document(BytesIO(content))
        paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        return "\n\n".join(paragraphs)
    if suffix in {".txt", ".md", ".do"}:
        return content.decode("utf-8")
    if suffix == ".pptx":
        return ""
    if suffix == ".mp4":
        return ""
    return ""


def _build_material_summary(arm_id: str | None, title: str, kind: str, extracted_text: str) -> str:
    arm_summary = ARM_FRAME_SUMMARIES.get(arm_id or "", "")
    excerpt = _truncate(extracted_text, 260) if extracted_text.strip() else ""
    if kind == "video":
        return f"{arm_summary} 该实验臂使用视频形式呈现材料，仿真当前使用该摘要作为视频内容的文本替身。"
    if arm_summary and excerpt:
        return f"{arm_summary} 材料摘录：{excerpt}"
    if arm_summary:
        return arm_summary
    if excerpt:
        return excerpt
    return title


def _copy_material_to_package(source_path: Path, package_dir: Path, relative_dir: Path) -> Path:
    target_dir = package_dir / relative_dir
    target_dir.mkdir(parents=True, exist_ok=True)
    target_path = target_dir / source_path.name
    shutil.copy2(source_path, target_path)
    return target_path


def _iter_sheet_rows(path: Path) -> tuple[list[str], list[list[Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise ValueError(f"Workbook {path} is empty")
    header_row = rows[0]
    headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
    data_rows = []
    for row in rows[1:]:
        if not any(cell is not None and str(cell).strip() for cell in row):
            continue
        padded = list(row) + [None] * (len(headers) - len(row))
        data_rows.append(padded[: len(headers)])
    return headers, data_rows


def _find_metric_column(headers: list[str], patterns: list[list[str]]) -> str | None:
    normalized_headers = {header: _normalize_text(header) for header in headers}
    for pattern_group in patterns:
        normalized_pattern = [_normalize_text(part) for part in pattern_group]
        for header, normalized in normalized_headers.items():
            if all(part in normalized for part in normalized_pattern):
                return header
    return None


def _numeric_value(value: Any) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    if text.endswith("%") and NUMERIC_PATTERN.match(text[:-1]):
        return float(text[:-1]) / 100
    if NUMERIC_PATTERN.match(text):
        return float(text)
    return None


def _average(values: list[float]) -> float:
    return round(sum(values) / len(values), 4)


def _aggregate_benchmarks(headers: list[str], rows: list[list[Any]]) -> tuple[dict[str, Any], dict[str, str]]:
    metrics: dict[str, Any] = {}
    metric_columns: dict[str, str] = {}
    header_index = {header: index for index, header in enumerate(headers)}
    for metric_def in XIHU_METRIC_DEFINITIONS:
        column = _find_metric_column(headers, metric_def["patterns"])
        if column is None:
            continue
        metric_columns[metric_def["key"]] = column
        values = [
            number
            for number in [_numeric_value(row[header_index[column]]) for row in rows]
            if number is not None
        ]
        if not values:
            continue
        metrics[metric_def["key"]] = {
            "label": metric_def["label"],
            "column": column,
            "mean": _average(values),
            "count": len(values),
        }

    family_metrics = [
        metrics[key]["mean"]
        for key in (
            "parents_enrollment_willingness",
            "children_enrollment_willingness",
            "spouse_enrollment_willingness",
        )
        if key in metrics
    ]
    if family_metrics:
        metrics["family_enrollment_willingness"] = {
            "label": "家庭投保意愿",
            "column": "derived",
            "mean": _average(family_metrics),
            "count": len(family_metrics),
        }
    return metrics, metric_columns


def import_xihu_round1(source: str, package_id: str = DEFAULT_XIHU_PACKAGE_ID) -> dict[str, Any]:
    source_root = Path(source)
    package_dir = get_xihu_package_dir(package_id)
    if package_dir.exists():
        shutil.rmtree(package_dir)
    package_dir.mkdir(parents=True, exist_ok=True)

    arm_directories = {arm_id: _match_single_path(source_root, arm_id) for arm_id in ARM_IDS}
    survey_root = source_root / "第一轮线下问卷数据"
    analysis_files = [
        path
        for path in _visible_files(source_root)
        if path.parent == source_root and path.suffix.lower() in {".docx", ".pptx", ".do"}
    ]

    materials: list[dict[str, Any]] = []
    arms: list[dict[str, Any]] = []
    survey_schema_arms: list[dict[str, Any]] = []
    benchmarks: dict[str, Any] = {}

    for arm_id in ARM_IDS:
        folder = arm_directories[arm_id]
        arm_materials: list[dict[str, Any]] = []
        for file_index, source_path in enumerate(_visible_files(folder), start=1):
            kind = _material_kind_for_file(arm_id, source_path)
            stored_path = _copy_material_to_package(source_path, package_dir, Path("materials") / arm_id)
            extracted_text = _extract_material_text(source_path)
            material = {
                "id": f"{arm_id}-{file_index:02d}-{_slugify(source_path.stem)}",
                "armId": arm_id,
                "armLabel": ARM_LABELS[arm_id],
                "kind": kind,
                "kindLabel": MATERIAL_KIND_LABELS[kind],
                "filename": source_path.name,
                "displayTitle": source_path.stem.replace("-", " ").replace("_", " "),
                "storageRef": str(stored_path.relative_to(package_dir)).replace("\\", "/"),
                "originalPath": str(source_path),
                "sha256": _sha256(source_path),
                "textSummary": _build_material_summary(arm_id, source_path.stem, kind, extracted_text),
                "textContent": extracted_text[:6000],
            }
            materials.append(material)
            arm_materials.append(material)

        workbook_path = survey_root / f"{arm_id}问卷.xlsx"
        headers, rows = _iter_sheet_rows(workbook_path)
        metric_map, metric_columns = _aggregate_benchmarks(headers, rows)
        survey_schema_arms.append(
            {
                "armId": arm_id,
                "armLabel": ARM_LABELS[arm_id],
                "headers": headers,
                "metricColumns": metric_columns,
            }
        )
        benchmarks[arm_id] = {
            "armId": arm_id,
            "armLabel": ARM_LABELS[arm_id],
            "sampleSize": len(rows),
            "metrics": metric_map,
        }
        arms.append(
            {
                "armId": arm_id,
                "label": ARM_LABELS[arm_id],
                "frameSummary": ARM_FRAME_SUMMARIES[arm_id],
                "materialIds": [material["id"] for material in arm_materials],
            }
        )

    shared_materials: list[dict[str, Any]] = []
    for file_index, source_path in enumerate(analysis_files, start=1):
        kind = _material_kind_for_file(None, source_path)
        stored_path = _copy_material_to_package(source_path, package_dir, Path("materials") / "shared")
        extracted_text = _extract_material_text(source_path)
        material = {
            "id": f"shared-{file_index:02d}-{_slugify(source_path.stem)}",
            "armId": None,
            "armLabel": None,
            "kind": kind,
            "kindLabel": MATERIAL_KIND_LABELS[kind],
            "filename": source_path.name,
            "displayTitle": source_path.stem.replace("-", " ").replace("_", " "),
            "storageRef": str(stored_path.relative_to(package_dir)).replace("\\", "/"),
            "originalPath": str(source_path),
            "sha256": _sha256(source_path),
            "textSummary": _build_material_summary(None, source_path.stem, kind, extracted_text),
            "textContent": extracted_text[:6000],
        }
        materials.append(material)
        shared_materials.append(material)

    package = {
        "packageId": package_id,
        "round": 1,
        "title": "西湖益联保第一轮干预材料",
        "importedAt": datetime.now(timezone.utc).isoformat(),
        "arms": arms,
        "materials": materials,
        "surveySchema": {
            "metricDefinitions": [
                {"key": metric["key"], "label": metric["label"]}
                for metric in XIHU_METRIC_DEFINITIONS
            ],
            "arms": survey_schema_arms,
            "openEndedQuestions": ["77", "78", "79"],
        },
        "benchmarks": benchmarks,
        "sourceMeta": {
            "sourceRoot": str(source_root),
            "armDirectories": {arm_id: str(path) for arm_id, path in arm_directories.items()},
            "surveyRoot": str(survey_root),
            "sharedMaterialIds": [material["id"] for material in shared_materials],
        },
    }

    get_xihu_package_path(package_id).write_text(
        json.dumps(package, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return package


def load_xihu_package(package_id: str) -> dict[str, Any]:
    package_path = get_xihu_package_path(package_id)
    if not package_path.exists():
        raise FileNotFoundError(f"Xihu package '{package_id}' not found at {package_path}")
    return json.loads(package_path.read_text(encoding="utf-8"))


def list_xihu_packages() -> list[dict[str, Any]]:
    root = ensure_xihu_package_root()
    packages = []
    for package_dir in sorted(path for path in root.iterdir() if path.is_dir()):
        package_path = package_dir / "package.json"
        if not package_path.exists():
            continue
        package = json.loads(package_path.read_text(encoding="utf-8"))
        packages.append(
            {
                "packageId": package["packageId"],
                "title": package.get("title", package["packageId"]),
                "round": package.get("round", 1),
                "importedAt": package.get("importedAt"),
            }
        )
    if packages:
        return packages
    return [{"packageId": DEFAULT_XIHU_PACKAGE_ID, "title": "西湖益联保第一轮干预材料", "round": 1}]


def get_xihu_package_options() -> list[dict[str, str]]:
    return [
        {
            "value": package["packageId"],
            "label": f"{package['packageId']} · 第{package['round']}轮",
        }
        for package in list_xihu_packages()
    ]


def get_xihu_arm_options() -> list[dict[str, str]]:
    return [{"value": arm_id, "label": ARM_LABELS[arm_id]} for arm_id in ARM_IDS]


def get_xihu_arm_label(arm_id: str) -> str:
    if arm_id not in ARM_LABELS:
        raise ValueError(f"Unsupported Xihu intervention arm: {arm_id}")
    return ARM_LABELS[arm_id]


def get_xihu_arm_summary(arm_id: str) -> str:
    if arm_id not in ARM_FRAME_SUMMARIES:
        raise ValueError(f"Unsupported Xihu intervention arm: {arm_id}")
    return ARM_FRAME_SUMMARIES[arm_id]


def normalize_xihu_arm_id(value: str) -> str:
    arm_value = str(value or "").strip()
    if not arm_value:
        raise ValueError("Xihu intervention arm is required")
    if arm_value in ARM_IDS:
        return arm_value

    candidate = arm_value[:2]
    if candidate in ARM_IDS:
        return candidate

    raise ValueError(f"Expected Xihu arm id in {ARM_IDS}, got '{arm_value}'")


def _knowledge_items_for_materials(package_id: str, arm_id: str, materials: list[dict[str, Any]]) -> list[dict[str, Any]]:
    summary_item = {
        "id": f"xihu-{package_id}-{arm_id}-summary",
        "title": f"{ARM_LABELS[arm_id]} 信息框架",
        "type": "text",
        "content": ARM_FRAME_SUMMARIES[arm_id],
        "enabled": True,
        "timestamp": datetime.now(timezone.utc).isoformat(),
    }
    material_items = [
        {
            "id": f"xihu-{package_id}-{material['id']}",
            "title": material["displayTitle"],
            "type": "text",
            "content": material["textSummary"],
            "enabled": True,
            "timestamp": datetime.now(timezone.utc).isoformat(),
        }
        for material in materials
    ]
    return [summary_item, *material_items]


def enrich_xihu_simulation_payload(
    scene_type: str,
    scene_config: dict[str, Any],
    agent_config: dict[str, Any] | None,
) -> tuple[dict[str, Any], dict[str, Any]]:
    if scene_type != "experiment_template":
        return scene_config, agent_config or {}

    enriched_scene_config, xihu_knowledge = enrich_xihu_scene_config(copy.deepcopy(scene_config))
    if not xihu_knowledge:
        return enriched_scene_config, agent_config or {}

    normalized_agent_config = copy.deepcopy(agent_config or {})
    enriched_agents = []
    for agent in normalized_agent_config.get("agents", []):
        existing_knowledge = list(agent.get("knowledgeBase") or agent.get("knowledge_base") or [])
        existing_ids = {item.get("id") for item in existing_knowledge}
        merged_knowledge = existing_knowledge + [
            item for item in xihu_knowledge if item["id"] not in existing_ids
        ]
        enriched_agent = copy.deepcopy(agent)
        enriched_agent["knowledgeBase"] = merged_knowledge
        enriched_agents.append(enriched_agent)

    normalized_agent_config["agents"] = enriched_agents
    return enriched_scene_config, normalized_agent_config


def enrich_xihu_scene_config(scene_config: dict[str, Any]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    generic_config = scene_config.get("generic_config")
    target_config = generic_config if generic_config is not None else scene_config
    parameters = dict(target_config.get("parameters") or {})
    scenario_id = str(target_config.get("scenario_id") or "")
    if scenario_id != "xihu_yilianbao":
        return scene_config, []

    package_id = str(parameters.get("xihu_package_id") or DEFAULT_XIHU_PACKAGE_ID)
    arm_id = normalize_xihu_arm_id(str(parameters.get("intervention_arm") or ""))
    package = load_xihu_package(package_id)

    selected_materials = [
        {
            **material,
            "downloadUrl": f"/api/xihu-round1/packages/{package_id}/materials/{material['id']}",
        }
        for material in package["materials"]
        if material.get("armId") == arm_id
    ]
    benchmark = package["benchmarks"][arm_id]
    material_brief = "\n".join(
        f"- {material['displayTitle']}: {material['textSummary']}" for material in selected_materials
    )

    enriched_parameters = {
        **parameters,
        "xihu_package_id": package_id,
        "intervention_arm": arm_id,
        "xihu_arm_label": ARM_LABELS[arm_id],
        "xihu_arm_summary": ARM_FRAME_SUMMARIES[arm_id],
        "xihu_material_brief": material_brief,
    }
    target_config["parameters"] = enriched_parameters
    target_config["xihu_package_id"] = package_id
    target_config["xihu_arm_id"] = arm_id
    target_config["xihu_arm_label"] = ARM_LABELS[arm_id]
    target_config["xihu_arm_summary"] = ARM_FRAME_SUMMARIES[arm_id]
    target_config["xihu_material_refs"] = selected_materials
    target_config["xihu_benchmarks"] = benchmark
    target_config["xihu_package_title"] = package["title"]
    target_config["xihu_provenance"] = package["sourceMeta"]

    return scene_config, _knowledge_items_for_materials(package_id, arm_id, selected_materials)
